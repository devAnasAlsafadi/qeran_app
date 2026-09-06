from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "artifacts" / "qa" / "Qeran_Jira_Manual_QA_Guide_2026-08-12.docx"
LOGO = ROOT / "assets" / "images" / "logo.png"

SKILL_SCRIPTS = Path(
    r"C:\Users\ACER\.codex\plugins\cache\openai-primary-runtime\documents\26.805.11740\skills\documents\scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry  # noqa: E402


# Qeran brand palette - sourced from lib/core/design_system/tokens/qeran_colors.dart
WINE = "431C33"
WINE_LIGHT = "4A1F38"
GOLD = "E4C094"
GOLD_DEEP = "B18454"
GOLD_LIGHT = "F2D9AC"
CREAM = "FBF4E6"
CANVAS = "F8F8F8"
PAPER = "FFFFFF"
INK = "201A1E"
MUTED = "6B6268"
SOFT_WINE = "F3EDF1"
SUCCESS = "256B4A"
FAIL = "9B1C1C"

FONT = "Arial"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 100, "bottom": 100, "start": 120, "end": 120}


CASES = [
    {
        "key": "QER-38",
        "title": "أيقونة التطبيق",
        "role": "مستخدم / خطّابة",
        "path": "شاشة الهاتف الرئيسية قبل فتح التطبيق",
        "pre": "تثبيت النسخة الحالية على جهاز Android وجهاز iPhone إن أمكن، ثم إزالة أي نسخة قديمة من الشاشة الرئيسية.",
        "steps": [
            "ثبّت التطبيق ثم ارجع إلى شاشة التطبيقات في الهاتف دون فتحه.",
            "افحص الأيقونة في الحجم العادي، ثم داخل قائمة التطبيقات وفي شاشة البحث.",
            "كرر الفحص على Android وiPhone وقارن الهوية بينهما.",
        ],
        "expected": [
            "لا توجد نقطة صفراء أسفل الرمز.",
            "اسم قِران والعلامة في المنتصف ولا يلامسان الحواف.",
            "الخاتم واللمعة واضحان، والهوية متقاربة على Android وiPhone.",
        ],
    },
    {
        "key": "QER-39",
        "title": "حركة الشعار عند التشغيل",
        "role": "مستخدم / خطّابة",
        "path": "تشغيل بارد للتطبيق - Splash",
        "pre": "إغلاق التطبيق بالكامل من التطبيقات الأخيرة. يفضّل الاختبار مرة قبل إكمال شاشات الترحيب ومرة بعدها.",
        "steps": [
            "أغلق التطبيق بالكامل ثم افتحه من الأيقونة.",
            "راقب حركة شعار قِران من أول لحظة حتى الانتقال إلى الترحيب أو الصفحة الرئيسية.",
            "كرر العملية على Android، ثم على iPhone خصوصًا عندما تكون شاشات الترحيب ما زالت مطلوبة.",
        ],
        "expected": [
            "الحركة تظهر في كل تشغيل بارد ولا تختفي على iPhone.",
            "الحركة واضحة وليست خاطفة على Android؛ مدتها المرئية تقارب 1.8 ثانية.",
            "بعدها يتم الانتقال للمسار الصحيح دون شاشة سوداء أو تعليق.",
        ],
    },
    {
        "key": "QER-14",
        "title": "تسمية حقل الاسم الظاهر",
        "role": "مستخدم",
        "path": "إنشاء حساب > بيانات الحساب، أو ملفي > عرض/تعديل الملف > تعديل الملف",
        "pre": "حساب جديد أو حساب يسمح بتعديل بيانات الملف.",
        "steps": [
            "افتح شاشة إدخال الاسم أثناء التسجيل.",
            "اقرأ عنوان الحقل والنص التوضيحي أسفله.",
            "بعد تسجيل الدخول افتح تعديل الملف وتأكد من بقاء التسمية نفسها.",
        ],
        "expected": [
            "اسم الحقل هو: الاسم الظاهر للمستخدمين.",
            "النص يوضح أن هذا هو الاسم الذي سيظهر داخل التطبيق، ولا يطلب اسمًا ثلاثيًا.",
        ],
    },
    {
        "key": "QER-1",
        "title": "إضافة وتعديل الصور بعد تخطيها في التسجيل",
        "role": "مستخدم",
        "path": "ملفي > عرض/تعديل الملف > عرض الملف الشخصي > إدارة الصور",
        "pre": "حساب تم إنشاؤه مع اختيار تخطي في خطوة إضافة الصور.",
        "steps": [
            "أكمل التسجيل بدون صور باستخدام تخطي.",
            "بعد الدخول افتح المسار المذكور واضغط إدارة الصور.",
            "اضغط إضافة صورة واختر صورة واضحة من الجهاز.",
            "أضف صورة ثانية، اجعلها رئيسية، ثم احذف إحدى الصور بعد تأكيد الحذف.",
        ],
        "expected": [
            "يمكن إضافة الصور حتى لو تم تخطيها أثناء التسجيل.",
            "الصورة الجديدة تظهر في القائمة مع حالة المعالجة/الاعتماد المناسبة.",
            "تغيير الصورة الرئيسية والحذف يعملان وتظهر رسالة نجاح لكل عملية.",
        ],
    },
    {
        "key": "QER-77",
        "title": "معاينة الصورة المرفوعة بالحجم الكامل",
        "role": "مستخدم",
        "path": "إضافة الصور أثناء التسجيل، أو ملفي > عرض/تعديل الملف > إدارة الصور",
        "pre": "وجود صلاحية وصول للصور وصورتين مختلفتين للاختبار.",
        "steps": [
            "أضف صورة وانتظر ظهورها داخل المربع المصغّر.",
            "اضغط على الصورة نفسها وليس علامة الحذف.",
            "أغلق المعاينة بزر الإغلاق أو الرجوع.",
            "أضف صورة ثانية وكرر المعاينة لكل صورة.",
        ],
        "expected": [
            "الضغط على أي صورة يفتح معاينة Full Screen للصورة الصحيحة.",
            "الإغلاق يعيدك لنفس شاشة الصور دون فقدان الاختيارات.",
            "علامة الحذف تبقى على الصورة المصغرة وتعمل بصورة مستقلة.",
        ],
    },
    {
        "key": "QER-15",
        "title": "لغة الإشعارات تتبع لغة التطبيق",
        "role": "مستخدم",
        "path": "ملفي > اللغة، ثم إشعار Push خارج التطبيق",
        "pre": "تفعيل الإشعارات، وجود حساب خطّابة أو مستخدم ثانٍ قادر على توليد إشعار، وإمكانية وضع التطبيق في الخلفية.",
        "steps": [
            "غيّر لغة التطبيق إلى العربية ثم ضعه في الخلفية.",
            "من الحساب الآخر أرسل رسالة أو نفّذ حدثًا يولد إشعارًا.",
            "افحص عنوان الإشعار ونصه في إشعارات الهاتف.",
            "غيّر لغة التطبيق إلى English وكرر الحدث فورًا دون إعادة تثبيت التطبيق.",
        ],
        "expected": [
            "الإشعار العربي يصل بالعربية عندما لغة التطبيق عربية.",
            "الإشعار التالي يصل بالإنجليزية مباشرة بعد تغيير اللغة إلى English.",
            "لا تبقى لغة الجهاز مسجلة على اللغة السابقة في السيرفر.",
        ],
    },
    {
        "key": "QER-19",
        "title": "رسالة الخطّابة الترحيبية تتبع لغة التطبيق",
        "role": "مستخدم",
        "path": "ملفي > اللغة، ثم رسائل > محادثة الخطّابة",
        "pre": "حساب مرتبط بخطّابة ويفضّل حساب جديد يمكن أن يستقبل رسالة الترحيب لأول مرة.",
        "steps": [
            "اختر English من إعدادات اللغة.",
            "افتح تبويب Messages ثم محادثة الخطّابة أو فعّل رسالة الترحيب بحساب جديد.",
            "افحص نص رسالة الترحيب واتجاه عرض النص.",
            "كرر على العربية للمقارنة.",
        ],
        "expected": [
            "رسالة الترحيب تكون بالإنجليزية عندما لغة التطبيق English.",
            "عند العربية تظهر بالعربية، واتجاه النص والواجهة مناسب للغة الحالية.",
        ],
    },
    {
        "key": "QER-58",
        "title": "توضيح وعمل زر استعادة المشتريات",
        "role": "مستخدم",
        "path": "ملفي > اشتراكي > عرض الباقات",
        "pre": "للاختبار الكامل استخدم حساب متجر سبق أن اشترى اشتراكًا، واستخدم حساب متجر آخر بلا مشتريات لاختبار الحالة السلبية.",
        "steps": [
            "افتح شاشة الباقات وانزل إلى زر استعادة المشتريات.",
            "تأكد من ظهور النص التوضيحي: إذا اشتريت اشتراكًا سابقًا على هذا الحساب.",
            "على حساب بلا مشتريات اضغط الزر وسجّل الرسالة.",
            "على الجهاز/حساب المتجر الذي لديه شراء سابق اضغط الزر وانتظر اكتمال المزامنة.",
        ],
        "expected": [
            "الزر واضح الغرض وليس بديلًا لشراء اشتراك جديد.",
            "الحساب بلا مشتريات يرى: لا توجد مشتريات سابقة للاستعادة.",
            "الحساب ذو الشراء السابق يستعيد الاستحقاق وتُحدّث حالة الاشتراك دون شراء جديد.",
        ],
    },
    {
        "key": "QER-5",
        "title": "البحث داخل القوائم الطويلة في الفلتر",
        "role": "مستخدم",
        "path": "زواج > زر الفلترة > الجنسية / بلد الإقامة / الوظيفة",
        "pre": "اتصال إنترنت وقائمة فلاتر محمّلة من السيرفر.",
        "steps": [
            "افتح الفلترة ثم افتح قائمة الجنسية.",
            "اكتب أول حرفين أو ثلاثة من قيمة موجودة واختر النتيجة.",
            "كرر على بلد الإقامة والوظيفة.",
            "ابحث بكلمة غير موجودة ثم امسحها.",
        ],
        "expected": [
            "يظهر حقل ابحث في الخيارات داخل كل قائمة طويلة.",
            "النتائج تتقلص أثناء الكتابة ويمكن اختيار القيمة الصحيحة.",
            "عند عدم وجود نتيجة تظهر رسالة لا توجد نتائج مطابقة دون تعطل القائمة.",
        ],
    },
    {
        "key": "QER-11",
        "title": "رسالة تعريف الفلتر لأول مرة",
        "role": "مستخدم",
        "path": "أول دخول إلى زواج بعد التسجيل وظهور الاقتراحات",
        "pre": "حساب جديد أو مسح بيانات التطبيق لإعادة حالة أول استخدام.",
        "steps": [
            "أكمل التسجيل حتى تظهر أول بطاقة في شاشة زواج.",
            "راقب الرسالة التعريفية واضغط حسنًا.",
            "انتقل إلى تبويب آخر ثم ارجع إلى زواج.",
            "أغلق التطبيق وافتحه مرة أخرى وادخل زواج.",
        ],
        "expected": [
            "تظهر رسالة مرحبًا بك مرة واحدة وتشرح استخدام زر بحث التصنيفات أعلى الصفحة.",
            "بعد إغلاقها لا تتكرر في نفس الحساب عند الرجوع أو إعادة تشغيل التطبيق.",
        ],
    },
    {
        "key": "QER-12",
        "title": "وضوح زر تخطي",
        "role": "مستخدم",
        "path": "زواج > شريط الأزرار أسفل البطاقة",
        "pre": "وجود ملف شخصي ظاهر في شاشة زواج.",
        "steps": [
            "افتح شاشة زواج وراقب زر تخطي وعلامة الإغلاق.",
            "قارن لون النص/الأيقونة مع الخلفية ومع زر التراجع.",
            "اضغط تخطي وتأكد أن الاستجابة البصرية لا تجعل الزر باهتًا.",
        ],
        "expected": [
            "زر تخطي بلون نبيذي داكن وواضح فوق خلفيته الفاتحة.",
            "حالة الضغط مفهومة ولا يختلط الزر مع الخلفية.",
        ],
    },
    {
        "key": "QER-24",
        "title": "تثبيت شريط التنقل وإخفاؤه أثناء التصفح",
        "role": "مستخدم",
        "path": "الصفحات الرئيسية، وبشكل خاص زواج",
        "pre": "بطاقة أو صفحة فيها محتوى قابل للتمرير عموديًا.",
        "steps": [
            "افتح كل تبويب رئيسي وتأكد أن شريط التنقل مثبت أسفل الشاشة ولا يغطي آخر نص.",
            "في زواج مرّر المحتوى إلى الأسفل حتى يظهر المزيد من معلومات البطاقة.",
            "راقب اختفاء الشريط أثناء النزول، ثم مرّر للأعلى.",
        ],
        "expected": [
            "الشريط ثابت في مكانه الطبيعي أسفل الصفحة ولا يبدو عائمًا فوق المحتوى.",
            "عند النزول في زواج يختفي لتوسيع مساحة القراءة، وعند الرجوع للأعلى يظهر من جديد.",
        ],
    },
    {
        "key": "QER-25",
        "title": "إظهار الصورة الحقيقية خلف التعتيم",
        "role": "مستخدم",
        "path": "زواج > صورة بطاقة الاقتراح",
        "pre": "وجود ملفين أو أكثر بصور مختلفة لم تُكشف للمستخدم بعد.",
        "steps": [
            "افتح بطاقة الملف الأول وراقب ألوان وأشكال الصورة المعتمة.",
            "انتقل إلى ملف ثانٍ له صورة مختلفة وقارن الخلفية.",
            "افتح الملف الكامل وتأكد من استمرار حماية الصورة دون تحولها إلى لون موحد.",
        ],
        "expected": [
            "تظهر الصورة الحقيقية خلف Blur أخف، ويمكن إدراك وجود صورة.",
            "لا تبدو كل الملفات كخلفية بنفسجية موحّدة.",
            "ملامح الشخص تبقى محمية ولا تصبح الصورة مكشوفة قبل السماح.",
        ],
    },
    {
        "key": "QER-29",
        "title": "ثبات مكان الإعجاب والتخطي في العربية والإنجليزية",
        "role": "مستخدم",
        "path": "زواج > أزرار البطاقة مع تبديل اللغة",
        "pre": "وجود اقتراحات والقدرة على تغيير لغة التطبيق.",
        "steps": [
            "اختر العربية وافتح زواج وسجّل مكان الإعجاب والتخطي.",
            "غيّر اللغة إلى English وافتح الشاشة نفسها.",
            "اسحب البطاقة يمينًا ثم استخدم زر الإعجاب في ملف آخر.",
            "اسحب بطاقة أخرى يسارًا ثم قارن مع زر التخطي.",
        ],
        "expected": [
            "الإعجاب يبقى في الجهة اليمنى والتخطي في الجهة اليسرى في اللغتين.",
            "السحب يمينًا يطابق الإعجاب والسحب يسارًا يطابق التخطي.",
        ],
    },
    {
        "key": "QER-41",
        "title": "تكبير أزرار التفاعل والتنقل",
        "role": "مستخدم",
        "path": "زواج > أزرار الإعجاب/التراجع/التخطي، وشريط التنقل الرئيسي",
        "pre": "وجود ملف ظاهر ومقارنة بصرية مع النسخة القديمة إن كانت متاحة.",
        "steps": [
            "افحص زر الإعجاب وقارنه بزرّي التراجع والتخطي.",
            "اضغط كل زر منفردًا وراقب الحركة والاستجابة.",
            "انتقل بين التبويبات وراقب حجم القرص الذهبي النشط وأيقونات القائمة.",
        ],
        "expected": [
            "الإعجاب هو الأكبر تقريبًا 72dp، والتراجع والتخطي يقاربان 60dp.",
            "أزرار القائمة الرئيسية أكبر وأكثر وضوحًا دون قص أو تداخل.",
            "ضغط زر واحد لا يسبب وميضًا أو حركة في بقية الأزرار.",
        ],
    },
    {
        "key": "QER-42",
        "title": "نقطة الرسائل غير المقروءة على القائمة الرئيسية",
        "role": "مستخدم",
        "path": "شريط التنقل الرئيسي > رسائل",
        "pre": "حساب مرتبط بخطّابة، وحساب الخطّابة قادر على إرسال رسالة جديدة بينما المستخدم في تبويب آخر.",
        "steps": [
            "ابقَ في تبويب زواج أو ملفي.",
            "من حساب الخطّابة أرسل رسالة جديدة للمستخدم.",
            "راقب أيقونة رسائل في شريط التنقل دون فتحها.",
            "افتح الرسائل واقرأ المحادثة ثم ارجع للشريط بعد تحديثها.",
        ],
        "expected": [
            "تظهر نقطة حمراء صغيرة فوق أيقونة رسائل عند وجود رسالة غير مقروءة.",
            "النقطة لا تظهر فوق تبويب آخر.",
            "بعد قراءة الرسائل وتحديث العدد تختفي النقطة.",
        ],
    },
    {
        "key": "QER-43",
        "title": "أيقونة زواج ومحاذاة التبويب النشط",
        "role": "مستخدم",
        "path": "شريط التنقل الرئيسي",
        "pre": "تسجيل دخول كمستخدم.",
        "steps": [
            "افتح الصفحة الرئيسية وافحص أيقونة زواج.",
            "اضغط إعجابات ثم رسائل ثم ملفي ثم ارجع إلى زواج.",
            "في كل انتقال راقب مكان القرص الذهبي بالنسبة لبقية الأيقونات.",
        ],
        "expected": [
            "أيقونة زواج تظهر كخاتم مع ماسة/رمز ماسي واضح.",
            "تلوين الأيقونة النشطة يبقى في نفس موضع الزر وعلى نفس خط الأيقونات، وليس مرتفعًا فوق القائمة.",
            "حركة الانتقال سلسة ولا تقفز للخلف عند الضغط السريع.",
        ],
    },
    {
        "key": "QER-44",
        "title": "إغلاق القوائم المنسدلة عند مسح الفلاتر",
        "role": "مستخدم",
        "path": "زواج > الفلترة",
        "pre": "قائمة طويلة مثل الجنسية أو بلد الإقامة متاحة.",
        "steps": [
            "افتح الفلترة واختر قيمة أو أكثر.",
            "افتح إحدى القوائم القابلة للبحث واتركها مفتوحة.",
            "اضغط مسح الكل.",
        ],
        "expected": [
            "تُمسح جميع الاختيارات.",
            "تُغلق القائمة المنسدلة المفتوحة فورًا ولا تبقى فوق شاشة الفلترة.",
            "يمكن اختيار قيم جديدة مباشرة دون إغلاق الشاشة يدويًا.",
        ],
    },
    {
        "key": "QER-46",
        "title": "إظهار بداية القسم التالي عندما تكون نبذة عني قصيرة",
        "role": "مستخدم",
        "path": "زواج > فتح الملف الشخصي الكامل",
        "pre": "ملف شخصي لديه نص قصير في نبذة عني ولديه قسم نبذة عن شريك حياتي أو معلومات لاحقة.",
        "steps": [
            "افتح الملف الكامل من بطاقة زواج.",
            "دون تمرير، راقب المساحة بعد نبذة عني.",
            "مرّر قليلًا للتأكد أن القسم الظاهر جزئيًا هو القسم التالي الحقيقي.",
        ],
        "expected": [
            "عند وجود مساحة كافية يظهر جزء من عنوان/محتوى القسم التالي أسفل نبذة عني.",
            "يفهم المستخدم بصريًا أن الصفحة تحتوي معلومات إضافية ويمكن تمريرها.",
            "لا توجد مساحة فارغة كبيرة تمنع ظهور القسم التالي.",
        ],
    },
    {
        "key": "QER-47",
        "title": "تصميم وألوان صفحة الملف الشخصي",
        "role": "مستخدم",
        "path": "زواج أو إعجابات > فتح الملف الشخصي الكامل",
        "pre": "ملف يحتوي عدة تصنيفات وأسئلة وإجابات واهتمامات.",
        "steps": [
            "افتح الملف الكامل ومرّر من أعلى الصفحة إلى أسفلها.",
            "افحص عناوين التصنيفات والأيقونات والخط الفاصل.",
            "قارن لون السؤال بلون الإجابة في أكثر من قسم.",
        ],
        "expected": [
            "عنوان التصنيف وأيقونته بهوية النبيذي، والأيقونة مناسبة للقسم.",
            "يوجد خط فاصل واضح بين عنوان التصنيف والمحتوى.",
            "الأسئلة بلون أسود/داكن والإجابات بالنبيذي، والبطاقات متناسقة مع تصميم غسان وهوية قِران.",
        ],
    },
    {
        "key": "QER-53",
        "title": "إعادة عرض الملفات بعد انتهاء الاقتراحات",
        "role": "مستخدم",
        "path": "زواج > نهاية مجموعة الملفات",
        "pre": "حساب يستطيع استعراض عدد محدود من الملفات حتى الوصول لنهاية القائمة، دون فلتر فعال.",
        "steps": [
            "استمر بالتخطي أو الإعجاب حتى تنتهي كل الملفات المحمّلة.",
            "انتظر اكتمال جلب الصفحات التالية إن وجدت.",
            "عند ظهور شاشة النهاية اضغط عرض الملفات مرة أخرى.",
            "كرر التجربة مع فلتر ضيق لا يعطي نتائج.",
        ],
        "expected": [
            "في نهاية القائمة غير المفلترة يظهر زر عرض الملفات مرة أخرى.",
            "الضغط يعيدك لأول ملف محمّل دون قتل التطبيق أو تسجيل دخول جديد.",
            "عند عدم وجود نتائج بسبب فلتر تظهر خيارات تعديل الفلترة وإزالتها بدل زر الإعادة.",
        ],
    },
    {
        "key": "QER-87",
        "title": "عرض أسئلة الفلترة الديناميكية وتطبيقها",
        "role": "مستخدم",
        "path": "زواج > زر الفلترة",
        "pre": "السيرفر يعيد أسئلة فلاتر من أنواع مدى، اختيار واحد، اختيار متعدد، اهتمامات، ونص.",
        "steps": [
            "افتح الفلترة وراجع أن الأسئلة المعادة من السيرفر كلها ظاهرة بعناوينها الصحيحة.",
            "حرّك مدى العمر/الطول إن وجد، واختر خيارًا مفردًا في سؤال آخر.",
            "اختر أكثر من قيمة في سؤال متعدد أو الاهتمامات، واكتب في حقل نصي/بحثي.",
            "اضغط عرض النتائج، ثم افتح الفلترة مرة أخرى.",
        ],
        "expected": [
            "كل نوع سؤال يظهر بأداة الإدخال المناسبة ولا يتم تجاهله أو تحويله لنوع خاطئ.",
            "الاختيارات تُرسل وتؤثر في النتائج، وتبقى محددة عند إعادة فتح الفلترة.",
            "مسح الكل يعيد كل الأسئلة لحالتها الافتراضية ويعيد النتائج غير المفلترة.",
        ],
    },
    {
        "key": "QER-54",
        "title": "إظهار زر الاستفسار كزر واضح في التوافق",
        "role": "مستخدم",
        "path": "إعجابات > التوافقات > بطاقة توافق بانتظار تبادل الصور",
        "pre": "وجود توافق في المرحلة التي يظهر فيها خيار أرسل استفساراتك للخطّابة.",
        "steps": [
            "افتح تبويب إعجابات ثم قسم التوافقات.",
            "ابحث عن بطاقة بانتظار تبادل الصور.",
            "افحص شكل خيار أرسل استفساراتك للخطّابة ثم اضغطه.",
        ],
        "expected": [
            "الخيار يظهر كزر حقيقي داكن/واضح بعرض مناسب وليس كنص عادي.",
            "حالة الضغط والتحميل ظاهرة، ولا يمكن إرسال الطلب مرتين أثناء التنفيذ.",
            "بعد النجاح ينتقل للمحادثة مع الخطّابة.",
        ],
    },
    {
        "key": "QER-55",
        "title": "صورة التوافق المعتمة بدل الدائرة البنفسجية",
        "role": "مستخدم",
        "path": "إعجابات > التوافقات",
        "pre": "وجود بطاقة توافق لشخص لديه صورة لكنها غير مكشوفة بعد.",
        "steps": [
            "افتح قائمة التوافقات وافحص صورة البطاقة قبل تبادل الصور.",
            "قارنها بصورة بطاقة أخرى ذات صورة مختلفة.",
            "انتقل بين بطاقات مراحل مختلفة إن توفرت.",
        ],
        "expected": [
            "تظهر الصورة الحقيقية بتمويه يحمي الهوية، لا دائرة بنفسجية مصمتة.",
            "يمكن تمييز اختلاف الخلفية بين الأشخاص مع بقاء الملامح محمية.",
        ],
    },
    {
        "key": "QER-56",
        "title": "الإعجاب أو التخطي من ملف مشارك داخل الدردشة",
        "role": "مستخدم",
        "path": "رسائل > محادثة الخطّابة > بطاقة ملف مشارك > عرض الملف الشخصي",
        "pre": "الخطّابة شاركت ملفين صالحين للمستخدم عبر المحادثة، ولم يسبق التفاعل معهما.",
        "steps": [
            "افتح البطاقة المشتركة الأولى واضغط عرض الملف الشخصي.",
            "تأكد من وجود زري إعجاب وتخطي مثبتين أسفل الملف.",
            "اضغط إعجاب في الملف الأول وانتظر رسالة النتيجة.",
            "افتح الملف الثاني واضغط تخطي.",
        ],
        "expected": [
            "زرّا إعجاب وتخطي يظهران فقط في سياق الملف المشارك من الدردشة.",
            "الإعجاب يرسل الطلب ثم يعيدك للمحادثة مع رسالة نجاح، أو يعرض Paywall إذا انتهى الحد.",
            "التخطي يعمل ويعيدك دون إرسال إعجاب.",
        ],
    },
    {
        "key": "QER-83",
        "title": "إرسال رسالة الخطوة الرسمية فورًا دون Refresh",
        "role": "مستخدم",
        "path": "إعجابات > التوافقات > توافق بعد تبادل الصور > خطوة رسمية عبر الخطّابة",
        "pre": "توافق وصل لمرحلة تسمح بالخطوة الرسمية، ومحادثة خطّابة متاحة.",
        "steps": [
            "افتح بطاقة التوافق المؤهلة واضغط خطوة رسمية عبر الخطّابة مرة واحدة.",
            "انتظر فتح محادثة الخطّابة دون تنفيذ Pull to refresh.",
            "افحص آخر رسالتين/العناصر في المحادثة.",
            "ارجع للتوافق وتأكد أن الزر لا يسمح بإرسال مكرر في نفس الحالة.",
        ],
        "expected": [
            "يظهر الملف الشخصي المشارك فورًا داخل المحادثة.",
            "تظهر معه فورًا رسالة: أرغب بالتقدّم رسمياً للخطوة التالية مع هذا الشخص.",
            "لا تحتاج المحادثة Refresh يدوي ولا تظهر نسخة مكررة بسبب الضغط المتكرر.",
        ],
    },
    {
        "key": "QER-85",
        "title": "مشاركة الملف مع رسالة الاستفسار قبل فتح المحادثة",
        "role": "مستخدم",
        "path": "إعجابات > التوافقات > أرسل استفساراتك للخطّابة",
        "pre": "بطاقة توافق يظهر فيها زر الاستفسار ومحادثة خطّابة متاحة.",
        "steps": [
            "اضغط أرسل استفساراتك للخطّابة على بطاقة محددة.",
            "عند فتح المحادثة لا تعمل Refresh.",
            "قارن الملف المشارك بالملف الذي ضغطت منه.",
            "اقرأ نص الرسالة المرافقة كاملًا.",
        ],
        "expected": [
            "يُشارك نفس الملف المطلوب الاستفسار عنه، وليس مجرد فتح الشات.",
            "تظهر فورًا الرسالة: السلام عليكم، أرغب بالاستفسار عن هذا الملف الشخصي ومعرفة مزيد من التفاصيل المتاحة عنه.",
            "يظهر الملف والنص دون تأخير أو تحديث يدوي.",
        ],
    },
    {
        "key": "QER-60",
        "title": "إظهار الإقامة والوظيفة والعمر في تطبيق الخطّابة",
        "role": "خطّابة",
        "path": "المستخدمون، المستخدمون > المشتركين > الإهتمامات، واستكشاف",
        "pre": "حساب خطّابة لديه مستخدمون تتضمن بياناتهم العمر وبلد الإقامة والوظيفة، مع مستخدم مشترك لديه اهتمامات.",
        "steps": [
            "افتح تبويب المستخدمون وافحص بطاقات انتظار المراجعة وغير المشتركين والمشتركين.",
            "في المشتركين اضغط أيقونة القلب/الإهتمامات وافحص بطاقات الوارد والصادر والتوافقات.",
            "افتح تبويب استكشاف وافحص عدة بطاقات.",
            "قارن البيانات مع الملف الكامل لنفس المستخدم.",
        ],
        "expected": [
            "تظهر بيانات العمر وبلد الإقامة والوظيفة عندما يرسلها السيرفر، دون قيم مختلقة.",
            "ترتيب الحقائق واضح ومختصر على البطاقات، ويمكن الوصول لبقية البيانات من عرض الملف.",
            "حالة الطلب/المؤقت تظهر أسفل معلومات المستخدم ولا تغطيها.",
        ],
    },
    {
        "key": "QER-61",
        "title": "عداد الوقت في صفحات اهتمامات الخطّابة",
        "role": "خطّابة",
        "path": "المستخدمون > المشتركين > أيقونة الإهتمامات",
        "pre": "مستخدم مشترك لديه إعجاب أو طلب نشط مع وقت انتهاء قادم من السيرفر.",
        "steps": [
            "افتح اهتمامات المستخدم وانتقل بين الوارد والصادر والتوافقات.",
            "ابحث عن بطاقة طلب نشط ودوّن الوقت المتبقي الظاهر.",
            "انتظر دقيقة أو أعد فتح الصفحة بعد فترة قصيرة.",
            "إن أمكن استخدم طلبًا قريب الانتهاء وراقب وصول العداد للصفر.",
        ],
        "expected": [
            "يظهر عداد وقت واضح على البطاقات التي لها مهلة.",
            "الوقت ينخفض بشكل صحيح ويعتمد على expiresAt/الوقت المتبقي من السيرفر.",
            "عند انتهاء المهلة تتحدّث الحالة ولا يبقى طلب منتهي قابلًا للتنفيذ.",
        ],
    },
    {
        "key": "QER-62",
        "title": "تحديث حالة التوافق من تطبيق الخطّابة",
        "role": "خطّابة",
        "path": "حالات التوافق > فتح حالة قابلة للتحديث > قسم الحالة",
        "pre": "حالة توافق فيها طلب رسمي وتسمح للخطّابة الحالية بالانتقال إلى حالة تالية. استخدم بيانات اختبار لا تؤثر على مستخدمين حقيقيين.",
        "steps": [
            "افتح الحالة وراجع المرحلة الحالية في رحلة التوافق.",
            "اضغط الإجراء الآمن التالي المتاح مثل تمت زيارة الأهل إذا ظهر.",
            "راقب حالة جاري التحديث ثم عودة الشاشة.",
            "افتح إجراء إغلاق/إلغاء إن ظهر وتأكد من ظهور نافذة التأكيد، ثم ألغِها ما لم تكن الحالة تجريبية.",
        ],
        "expected": [
            "تظهر فقط الانتقالات القانونية التي يسمح بها السيرفر للحالة الحالية.",
            "بعد التحديث تتغير الحالة والـTimeline وتظهر رسالة نجاح دون تكرار الطلب.",
            "الإجراءات النهائية/الخطرة تطلب تأكيدًا، والحالات غير القابلة للتحديث تعرض سببًا واضحًا بدل زر ميت.",
        ],
    },
    {
        "key": "QER-65",
        "title": "فلاتر حالات التوافق مطابقة للوحة التحكم",
        "role": "خطّابة",
        "path": "حالات التوافق > زر التصفية",
        "pre": "وجود حالات في أكثر من مرحلة ووجود طلبات رسمية نشطة وغير نشطة.",
        "steps": [
            "افتح فلتر الحالات وراجع خيارات الحالة.",
            "اختبر: تم قبول الإعجاب، بانتظار تبادل الصور، تم تبادل الصور، مرفوض، ومنتهي حسب المتاح.",
            "اختبر فلتر الطلب الرسمي: الكل، يوجد طلب رسمي نشط، لا يوجد طلب رسمي نشط.",
            "ادمج فلتر مرحلة مع فلتر طلب رسمي واضغط تطبيق، ثم جرّب مسح التصفية.",
        ],
        "expected": [
            "خيارات الفلتر تطابق معاني وقيم لوحة التحكم.",
            "الفلتران يعملان معًا على السيرفر وتظهر فقط الحالات المطابقة.",
            "مسح التصفية يعيد القائمة الكاملة، والحالة الفارغة تشرح عدم وجود نتائج مطابقة.",
        ],
    },
    {
        "key": "QER-69",
        "title": "مشاركة رسالة كود الإحالة كاملة",
        "role": "خطّابة",
        "path": "أي شاشة رئيسية للخطّابة > ترس الإعدادات أعلى الشاشة > حسابي > كود الإحالة",
        "pre": "حساب خطّابة لديه كود إحالة ظاهر.",
        "steps": [
            "افتح حسابي وانزل إلى بطاقة كود الإحالة.",
            "اضغط مشاركة الكود.",
            "في نافذة المشاركة اختر تطبيق ملاحظات/رسائل تجريبيًا أو انسخ المعاينة واقرأ النص.",
            "قارن الكود داخل الرسالة بالكود الظاهر في البطاقة.",
        ],
        "expected": [
            "تفتح نافذة المشاركة بنص كامل: حمّل تطبيق قِران واستفد من الخصم عند الاشتراك باستخدام كود الإحالة: [الكود].",
            "الكود داخل الرسالة صحيح وغير مقطوع، ويمكن نسخ الكود منفردًا أيضًا.",
        ],
    },
    {
        "key": "QER-72",
        "title": "إخفاء تعديل النص لملف غير مسند للخطّابة",
        "role": "خطّابة",
        "path": "استكشاف > فتح ملف غير مسند، ثم المستخدمون > فتح ملف مسند",
        "pre": "حساب خطّابة، وملفان: واحد مسند لها والآخر مسند لخطّابة أخرى/غير مسند.",
        "steps": [
            "من استكشاف افتح ملفًا غير موجود ضمن مستخدميك.",
            "مرّر أقسام نبذة عني ونبذة عن الشريك والإجابات النصية وابحث عن قلم التعديل.",
            "ارجع إلى المستخدمون وافتح ملفًا مسندًا لك.",
            "افحص الأقسام نفسها وحاول فتح تعديل نص تجريبي دون حفظ تغيير غير مرغوب.",
        ],
        "expected": [
            "لا يظهر قلم التعديل إطلاقًا في الملف غير المسند لهذه الخطّابة.",
            "يظهر التعديل فقط في الملف المسند لها وبحسب صلاحيات السيرفر.",
            "أي محاولة غير مصرح بها من السيرفر تُعامل كمنع صلاحية ولا تكشف زرًا مضللًا.",
        ],
    },
    {
        "key": "QER-75",
        "title": "التواصل مع الخطّابة المسؤولة من داخل الملف",
        "role": "خطّابة",
        "path": "المستخدمون أو استكشاف > فتح الملف الكامل > أسفل الصفحة",
        "pre": "ملف شخصي لديه خطّابة مسؤولة معروفة ومحادثة قابلة للفتح.",
        "steps": [
            "افتح الملف الكامل ومرّر حتى نهاية جميع المواصفات.",
            "ابحث عن خيار التواصل مع الخطّابة المسؤولة.",
            "اضغط الخيار وانتظر فتح وسيلة التواصل.",
        ],
        "expected": [
            "يظهر الخيار واضحًا في أسفل الملف بعد جميع المعلومات.",
            "الضغط يفتح المحادثة المباشرة أو بيانات التواصل مع الخطّابة المسؤولة عن هذا الشخص بالذات.",
            "لا يتم فتح محادثة مع خطّابة غير مرتبطة بالملف.",
        ],
    },
    {
        "key": "QER-76",
        "title": "فتح ملف المستخدم من رأس المحادثة",
        "role": "خطّابة",
        "path": "المحادثات > فتح محادثة مستخدم > الضغط على الاسم أو الصورة أعلى المحادثة",
        "pre": "وجود محادثة مستخدم في تطبيق الخطّابة.",
        "steps": [
            "افتح المحادثة وسجّل اسم وصورة المستخدم في رأس الشاشة.",
            "اضغط الصورة، ثم كرر التجربة بالضغط على الاسم إن أمكن.",
            "راجع المعلومات والصور المسموح للخطّابة برؤيتها.",
            "اضغط رجوع.",
        ],
        "expected": [
            "يفتح الملف الشخصي الكامل لنفس مستخدم المحادثة.",
            "تظهر البيانات والمواصفات وفق صلاحيات الخطّابة دون بحث يدوي.",
            "الرجوع يعيد إلى المحادثة نفسها وفي موضعها السابق.",
        ],
    },
]


PENDING = [
    ("QER-18", "تعديل اسم المستخدم", "To Do", "يحتاج مسار API وصيغة الطلب وقواعد الموافقة من Backend."),
    ("QER-31", "Google Login على iPhone", "In Progress", "يحتاج جهاز iPhone حقيقي وحساب Google لاختبار المسار قبل اعتماد النتيجة."),
    ("QER-50", "بيانات ناقصة في قائمة الاهتمامات", "To Do", "واجهة الموبايل جاهزة؛ يلزم التأكد أن Backend يعيد العمر والإقامة والوظيفة لهذه القائمة."),
    ("QER-52", "تأخر قبول الملف بعد الموافقة", "To Do", "يلزم تتبع حادثة فعلية بالوقت وuserId في السيرفر."),
    ("QER-57", "كود الخصم عند شراء باقة", "To Do", "Android مدعوم؛ iOS يحتاج توقيع StoreKit Promotional Offer من Backend أو آلية Offer Codes."),
    ("QER-73", "عدد حالات التوافق في الرئيسية", "Not Needed", "التذكرة مستبعدة رسميًا ولا تدخل ضمن فحص Done."),
]


GROUPS = [
    ("أولًا: التشغيل والتسجيل والحساب", ["QER-38", "QER-39", "QER-14", "QER-1", "QER-77", "QER-15", "QER-19", "QER-58"]),
    ("ثانيًا: شاشة زواج والفلترة والملف", ["QER-5", "QER-11", "QER-12", "QER-24", "QER-25", "QER-29", "QER-41", "QER-42", "QER-43", "QER-44", "QER-46", "QER-47", "QER-53", "QER-87"]),
    ("ثالثًا: الإعجابات والتوافقات والدردشة", ["QER-54", "QER-55", "QER-56", "QER-83", "QER-85"]),
    ("رابعًا: تطبيق الخطّابة", ["QER-60", "QER-61", "QER-62", "QER-65", "QER-69", "QER-72", "QER-75", "QER-76"]),
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size: float | None = None, color: str = INK, bold: bool | None = None, italic: bool | None = None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        tag = f"w:{edge_name}"
        tag_el = tc_borders.find(qn(tag))
        if tag_el is None:
            tag_el = OxmlElement(tag)
            tc_borders.append(tag_el)
        for key, value in edge_data.items():
            tag_el.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_cell_text(cell, size=9.0, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        set_rtl(paragraph, align)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc: Document, num_fmt: str, text: str, is_bullet: bool = False) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abs, default=0) + 1
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    lvl.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "right")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "540")
    ind.set(qn("w:hanging"), "260")
    p_pr.append(ind)
    bidi = OxmlElement("w:bidi")
    p_pr.append(bidi)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:cs"), FONT)
    r_pr.append(fonts)
    if is_bullet:
        color = OxmlElement("w:color")
        color.set(qn("w:val"), GOLD_DEEP)
        r_pr.append(color)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_list_item(doc: Document, text: str, num_id: int, bullet: bool = False):
    paragraph = doc.add_paragraph()
    set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    apply_num(paragraph, num_id)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(paragraph)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(
        run,
        size={1: 16, 2: 13, 3: 11.5}[level],
        color=WINE if level < 3 else GOLD_DEEP,
        bold=True,
    )
    return paragraph


def add_callout_paragraph(doc: Document, label: str, text: str, fill=CREAM, border=GOLD):
    paragraph = doc.add_paragraph()
    set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), border)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    r1 = paragraph.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=WINE, bold=True)
    r2 = paragraph.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return paragraph


def add_metadata_table(doc: Document, case: dict):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    rows = [("الدور", case["role"]), ("المسار", case["path"])]
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        set_cell_shading(table.cell(i, 0), GOLD_LIGHT)
        set_cell_shading(table.cell(i, 1), PAPER if i % 2 == 0 else CANVAS)
        style_cell_text(table.cell(i, 0), size=9.5, color=WINE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell_text(table.cell(i, 1), size=9.5, color=INK)
        prevent_row_split(table.rows[i])
    apply_table_geometry(
        table,
        [1728, 7632],
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS_DXA,
    )
    return table


def add_result_box(doc: Document):
    paragraph = doc.add_paragraph()
    set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), SOFT_WINE)
    p_pr.append(shd)
    r = paragraph.add_run("النتيجة الفعلية:  [ ] ناجح    [ ] فشل    [ ] متعذر بسبب البيانات     ملاحظات: ____________________")
    set_run_font(r, size=9.5, color=WINE, bold=True)
    return paragraph


def add_case(doc: Document, case: dict, decimal_num: int, bullet_num: int):
    heading = add_heading(doc, f"{case['key']} - {case['title']}", level=2)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)
    add_metadata_table(doc, case)
    add_callout_paragraph(doc, "قبل الاختبار", case["pre"])
    add_heading(doc, "خطوات الاختبار", level=3)
    for step in case["steps"]:
        add_list_item(doc, step, decimal_num)
    add_heading(doc, "النتيجة المتوقعة", level=3)
    for item in case["expected"]:
        add_list_item(doc, item, bullet_num, bullet=True)
    add_result_box(doc)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, before, after, color in [
        (1, 16, 18, 10, WINE),
        (2, 13, 14, 7, WINE),
        (3, 11.5, 10, 5, GOLD_DEEP),
    ]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("دليل الاختبار الحي  |  قِران")
    set_run_font(hr, size=8.5, color=WINE, bold=True)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), GOLD)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    even_header = section.even_page_header
    ehp = even_header.paragraphs[0]
    ehp.clear()
    ehp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ehr = ehp.add_run("دليل الاختبار الحي  |  قِران")
    set_run_font(ehr, size=8.5, color=WINE, bold=True)
    ehp_pr = ehp._p.get_or_add_pPr()
    ehp_bdr = OxmlElement("w:pBdr")
    even_bottom = OxmlElement("w:bottom")
    even_bottom.set(qn("w:val"), "single")
    even_bottom.set(qn("w:sz"), "6")
    even_bottom.set(qn("w:color"), GOLD)
    ehp_bdr.append(even_bottom)
    ehp_pr.append(ehp_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("قِران - دليل فحص تعديلات جيرا   |   صفحة ")
    set_run_font(fr, size=8.5, color=MUTED)
    add_field(fp, "PAGE")

    even_footer = section.even_page_footer
    efp = even_footer.paragraphs[0]
    efp.clear()
    efp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    efr = efp.add_run("قِران - دليل فحص تعديلات جيرا   |   صفحة ")
    set_run_font(efr, size=8.5, color=MUTED)
    add_field(efp, "PAGE")

    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.clear()
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = ffp.add_run("قِران  |  12 أغسطس 2026")
    set_run_font(rr, size=8.5, color=MUTED)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    run = p.add_run()
    run.add_picture(str(LOGO), width=Inches(1.55))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", "شعار تطبيق قِران")
        drawing[0].set("title", "قِران")

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(8)
    kr = kicker.add_run("دليل الاختبار الحي")
    set_run_font(kr, size=11, color=GOLD_DEEP, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    tr = title.add_run("دليل فحص تعديلات جيرا")
    set_run_font(tr, size=28, color=WINE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    sr = subtitle.add_run("تطبيق قِران - الموبايل وتطبيق الخطّابة")
    set_run_font(sr, size=14, color=WINE_LIGHT, bold=True)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    meta = [
        ("نطاق الدليل", "35 تذكرة مكتملة بحالة Done"),
        ("الأدوار", "المستخدم + الخطّابة"),
        ("تاريخ التجهيز", "12 أغسطس 2026"),
        ("نسخة التطبيق", "البناء الحالي للمشروع - اكتب رقم النسخة هنا: __________"),
    ]
    for i, (label, value) in enumerate(meta):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        set_cell_shading(table.cell(i, 0), GOLD_LIGHT)
        set_cell_shading(table.cell(i, 1), PAPER if i % 2 == 0 else CANVAS)
        style_cell_text(table.cell(i, 0), size=10, color=WINE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell_text(table.cell(i, 1), size=10, color=INK)
        prevent_row_split(table.rows[i])
    apply_table_geometry(table, [2304, 7056], table_width_dxa=CONTENT_WIDTH_DXA, indent_dxa=TABLE_INDENT_DXA, cell_margins_dxa=CELL_MARGINS_DXA)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(28)
    nr = note.add_run("الهدف: اختبار كل تعديل حيًا، وتسجيل ناجح / فشل / متعذر قبل فتح أي تذكرة جديدة.")
    set_run_font(nr, size=10.5, color=MUTED)
    doc.add_page_break()


def add_intro(doc: Document, bullet_num: int):
    add_heading(doc, "طريقة استخدام هذا الدليل", level=1)
    add_callout_paragraph(
        doc,
        "قاعدة النتيجة",
        "اعتبر الحالة ناجحة فقط عند تطابق النتيجة المرئية والسلوك معًا. إذا كانت بيانات السيرفر المطلوبة غير موجودة، اختر متعذر بسبب البيانات بدل تسجيل فشل على الموبايل.",
        fill=CREAM,
        border=GOLD_DEEP,
    )
    for item in [
        "اختبر على النسخة المبنية من الكود الحالي، وسجّل رقم النسخة أو تاريخ الـAPK في الغلاف.",
        "نفّذ التذاكر بالترتيب داخل كل رحلة حتى تقلل تبديل الحسابات.",
        "استخدم ناجح عندما تطابق كل النتائج، وفشل عند وجود فرق قابل للتكرار، ومتعذر عندما تنقص بيانات أو صلاحيات من السيرفر.",
        "عند الفشل سجّل الجهاز، نظام التشغيل، الحساب، الخطوات، وصورة أو فيديو قصير.",
    ]:
        add_list_item(doc, item, bullet_num, bullet=True)

    add_heading(doc, "حسابات وبيانات الاختبار المقترحة", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["الرمز", "الحساب/البيانات", "الاستخدام"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        set_cell_shading(table.cell(0, i), WINE)
        style_cell_text(table.cell(0, i), size=9.5, color=PAPER, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    test_data = [
        ("U1", "مستخدم مشترك وله صورة وخطّابة", "زواج، إعجابات، توافقات، محادثة"),
        ("U2", "مستخدم من الجنس المقابل وله صورة", "إنشاء إعجاب/توافق مع U1"),
        ("U3", "مستخدم جديد تخطى الصور", "QER-1 وQER-11 وQER-77"),
        ("M1", "خطّابة مسؤولة عن U1", "اختبارات تطبيق الخطّابة والرسائل"),
        ("M2", "خطّابة مسؤولة عن U2", "التواصل بين الخطّابات وصلاحيات الملفات"),
        ("متاجر", "حساب متجر بمشتريات وآخر بلا مشتريات", "QER-58"),
    ]
    for idx, row_data in enumerate(test_data, start=1):
        row = table.add_row()
        for col, value in enumerate(row_data):
            row.cells[col].text = value
            set_cell_shading(row.cells[col], PAPER if idx % 2 else CANVAS)
            style_cell_text(row.cells[col], size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
        prevent_row_split(row)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [900, 3600, 4860], table_width_dxa=CONTENT_WIDTH_DXA, indent_dxa=TABLE_INDENT_DXA, cell_margins_dxa=CELL_MARGINS_DXA)

    add_heading(doc, "خط الأساس التقني قبل الاختبار الحي", level=1)
    add_callout_paragraph(
        doc,
        "آخر تحقق",
        "Flutter Analyze بلا ملاحظات، 930 من 930 اختبار ناجح، وبناء Debug APK اكتمل بنجاح. هذا لا يغني عن اختبار السيرفر والإشعارات والمتاجر والأجهزة الحقيقية في هذا الدليل.",
        fill=SOFT_WINE,
        border=WINE,
    )


def add_summary_table(doc: Document):
    doc.add_page_break()
    add_heading(doc, "جدول التغطية السريع - جميع التعديلات المكتملة", level=1)
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("استخدم هذا الجدول كلوحة متابعة، ثم ارجع إلى حالة الاختبار التفصيلية بالأسفل.")
    set_run_font(r, size=10.5, color=MUTED)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["#", "جيرا", "الدور", "الشاشة/المسار", "التعديل المختصر", "النتيجة"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        set_cell_shading(table.cell(0, i), WINE)
        style_cell_text(table.cell(0, i), size=8.5, color=PAPER, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    case_by_key = {c["key"]: c for c in CASES}
    ordered_keys = [key for _, keys in GROUPS for key in keys]
    for idx, key in enumerate(ordered_keys, start=1):
        case = case_by_key[key]
        row = table.add_row()
        vals = [str(idx), key, case["role"], case["path"], case["title"], "[ ]"]
        for col, value in enumerate(vals):
            row.cells[col].text = value
            set_cell_shading(row.cells[col], PAPER if idx % 2 else CANVAS)
            style_cell_text(row.cells[col], size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER if col in (0, 1, 2, 5) else WD_ALIGN_PARAGRAPH.RIGHT)
        prevent_row_split(row)
    apply_table_geometry(
        table,
        [420, 900, 1050, 2100, 4170, 720],
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=90,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 90, "end": 90},
    )


def add_detailed_cases(doc: Document, decimal_num: int, bullet_num: int):
    case_by_key = {c["key"]: c for c in CASES}
    for group_title, keys in GROUPS:
        doc.add_page_break()
        add_heading(doc, group_title, level=1)
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(8)
        rr = p.add_run(f"عدد حالات الاختبار في هذا القسم: {len(keys)}")
        set_run_font(rr, size=9.5, color=MUTED)
        for key in keys:
            add_case(doc, case_by_key[key], decimal_num, bullet_num)


def add_pending_appendix(doc: Document):
    doc.add_page_break()
    add_heading(doc, "ملحق: تذاكر لم تُحسم بعد", level=1)
    add_callout_paragraph(
        doc,
        "مهم",
        "هذه التذاكر لم تُغلق ضمن دفعة الموبايل. لا تسجلها فشلًا من هذا الدليل؛ ارجع لسبب الانتظار المحدد لكل تذكرة.",
        fill=CREAM,
        border=GOLD_DEEP,
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["جيرا", "العنوان", "الحالة", "سبب عدم الحسم"]):
        table.cell(0, i).text = h
        set_cell_shading(table.cell(0, i), WINE)
        style_cell_text(table.cell(0, i), size=9, color=PAPER, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for idx, data in enumerate(PENDING, start=1):
        row = table.add_row()
        for col, value in enumerate(data):
            row.cells[col].text = value
            set_cell_shading(row.cells[col], PAPER if idx % 2 else CANVAS)
            style_cell_text(row.cells[col], size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER if col in (0, 2) else WD_ALIGN_PARAGRAPH.RIGHT)
        prevent_row_split(row)
    apply_table_geometry(table, [1050, 2450, 1250, 4610], table_width_dxa=CONTENT_WIDTH_DXA, indent_dxa=TABLE_INDENT_DXA, cell_margins_dxa=CELL_MARGINS_DXA)

    add_heading(doc, "ملخص التنفيذ بعد الجولة", level=1)
    for label in [
        "عدد الحالات الناجحة: ______ من 35",
        "عدد الحالات الفاشلة: ______",
        "عدد الحالات المتعذرة بسبب البيانات/السيرفر: ______",
        "روابط الصور أو الفيديوهات: ______________________________________________",
        "ملاحظات عامة: _________________________________________________________",
    ]:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(9)
        r = p.add_run(label)
        set_run_font(r, size=10.5, color=INK)


def build():
    assert len(CASES) == 35, f"Expected 35 cases, got {len(CASES)}"
    actual_keys = {c["key"] for c in CASES}
    grouped_keys = [k for _, keys in GROUPS for k in keys]
    assert len(grouped_keys) == 35 and set(grouped_keys) == actual_keys

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    decimal_num = add_numbering_definition(doc, "decimal", "%1.")
    bullet_num = add_numbering_definition(doc, "bullet", "●", is_bullet=True)

    add_cover(doc)
    add_intro(doc, bullet_num)
    add_summary_table(doc)
    add_detailed_cases(doc, decimal_num, bullet_num)
    add_pending_appendix(doc)

    core = doc.core_properties
    core.title = "دليل فحص تعديلات جيرا - تطبيق قِران"
    core.subject = "دليل اختبار حي لـ35 تعديلًا مكتملًا على الموبايل وتطبيق الخطّابة"
    core.author = "Qeran Mobile Team"
    core.keywords = "Qeran, Jira, QA, Flutter, Mobile, Matchmaker"
    core.comments = "Generated for manual live verification on 2026-08-12"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
